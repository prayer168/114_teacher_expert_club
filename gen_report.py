# -*- coding: utf-8 -*-
"""產生 114 學年度南科綠手指教師專業學習社群成果報告 (docx)。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

IMG = "/tmp/claude-0/-home-user-114-teacher-expert-club/4dac3fae-3612-5b48-8ea2-2e48125ab8a6/scratchpad/imgs"
OUT = "/home/user/114_teacher_expert_club/114南科綠手指教師專業社群成果報告.docx"

FONT = "標楷體"

def set_cjk(run, size=None, bold=None, color=None):
    run.font.name = FONT
    r = run._element
    rpr = r.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for a in ('w:eastAsia', 'w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(a), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color

def para(doc, text="", size=12, bold=False, align=None, color=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if text:
        run = p.add_run(text)
        set_cjk(run, size=size, bold=bold, color=color)
    return p

def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def cell_text(cell, text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run(text)
    set_cjk(run, size=size, bold=bold)
    if fill:
        shade_cell(cell, fill)

def add_heading_bar(doc, text):
    p = para(doc, text, size=14, bold=True, space_after=4)
    # 底線效果用段落框線
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '2'); bottom.set(qn('w:color'), '2E7D32')
    pbdr.append(bottom)
    pPr.append(pbdr)
    p.paragraph_format.space_before = Pt(10)
    return p

HDR = "D9EAD3"   # 淺綠標題底

def info_table(doc, rows):
    """rows: list of (label,value) or (label,value,label2,value2)"""
    t = doc.add_table(rows=0, cols=4)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in rows:
        tr = t.add_row().cells
        if len(row) == 2:
            cell_text(tr[0], row[0], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill=HDR)
            a = tr[1].merge(tr[2]).merge(tr[3])
            cell_text(a, row[1])
        else:
            cell_text(tr[0], row[0], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill=HDR)
            cell_text(tr[1], row[1])
            cell_text(tr[2], row[2], bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill=HDR)
            cell_text(tr[3], row[3])
    return t

def photo_grid(doc, ids, cols=3, width_cm=5.0):
    ids = [i for i in ids if os.path.exists(os.path.join(IMG, i + ".jpg"))]
    if not ids:
        return
    rows = (len(ids) + cols - 1) // cols
    t = doc.add_table(rows=rows, cols=cols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for idx, fid in enumerate(ids):
        r, c = divmod(idx, cols)
        cell = t.cell(r, c)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(os.path.join(IMG, fid + ".jpg"), width=Cm(width_cm))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ---------------- 影像對應 ----------------
A1 = "1nFq4fSJT1Nj23esXPhlViieq3nAUn2HB 1VEzgi48gxl7QhiDJI4LblxuSziIdsO-b 1Y8WbbrGVWZ_V2ucRGjq6Q99_eDLH01PM 1ySvLdladVgCIZF6d6wkrOBx_toPT9QDR 1RXvpPXpefZP0fYmmFJIYKtorpxgEeFS2 1abYP8dXYchxml5WudaNbn4iOzTb0aH-M".split()
A2 = "1k7lKSWGbh7jnFlBc5G_ctni_XHhKy5e9 1LFKu6MLYX3Nlk1gy06VehZFEa_9MDZxf 1Jp4ZB9CWRhMBD1oFeRLfI9PJExXKNsW8 1kVnFF2GhRzOPMVrtqOO9qSfPwgAnSBCZ 1a0kwUn7NG5mTpnaZwSdfG0F4VEoyGcUv 1j1MW7OR9DTlnHLgrJKL3gNuCHGTnP39s".split()
A3 = "1sCtiOeL-HHuHMS_kt04TD9T1t-hZkp-Q 1VNnZYXPa50F8J1NgyRtgrWHW7iRoLlQq 1-3hX_2jdX9DQTSRVIHu4oOSrXW-EprAL 1DjaNg8eYwzC9lx1yFmEpjTMkcKYKQgYT 1LHF103CMtG2ozBSW3-xNApGN3BwBlj1W 1tYkEH_djA30Egy20i7Xdm224TuChwnFq".split()
A4 = "1EUc4rUh0n4Z_fHVqnTcKqhjVfT3YiiiT 12Jfn2brJMXakd8Vp8oqb0L5e2VBoaukg 1fgMQxN_I0YX2pEuu4lQTcFXVPMmCTfoY 1bT1Cb52pQMw-xgmwhLhp34_vuixiMRL3 1kzH6mWDfaU6Q_7sp76qOAjiV1zCw5gFO 1fA9e6bI9dEIHfbIy8N3lmVjuIMTDGnen".split()
A5 = "1mFZQKgWOhkbBU4pCBKrdvG4Z45-s423j 1ZEbsl91ro5BoCKYiHafVwAZtBMYeu92p 1S8JE4G0PImGzVfuPRO43P3rbwy8Yn-0M 1pMnioiGf4p1MzAuljgaTlNvZRBJDpkO0 18djj9gTbHghGqiqyFTFaEmn5RVDxS65e 1gQGGzoOXVslf40E-7Q_z6cBWdO5gOtO7".split()
A6 = "1-IEbthVsZxctqCNWubGbnLuyFaUHWDWY 1eEXOrkivdUKrGIN94g5DI6ob8RK6gBYy".split()
R  = "1w1jEh5NY0WYk7MC5DV0Yz96P7iYnID2k 1-0jNAO0G27aXNjUt5EJPoa3R6j3WhnCi 1gSHMCs7oJVnzCqPBa3l_RlNn12F0hzg4 1GQzxzIac7flU0SHNP1RMSlk_lTDLc75P 1-EFPNZaHaQT9khZbedPVqBTJxnSNdQsc 1Jml9RG0QOXqcZ0wxQ_m31ok2CS919UkB".split()
P1 = "1fZzASVo2h7b7OCB9u2EFiCEi1Ambr_C8 1bGaoVrYOJSWXgoS1BkRXZYByLue-_vZQ 1r4asgjiMl_fkR1RYEbFJEFCXM3J8r3bC 1TRr8lEK4S4UQOaYPLk2vDQaYTylXUxvA 1lZBAV8oPeJpmxq1RvmBUk39qWOiVdn-m 10DNYZWa96PiFfIUTYvZY2vHB0_a_VsAF".split()
P2 = "1vW0JMnbkLenu34iDGPpz3ESNqk3wyUtH 1_ohrMatXiVCseK38AjL8XW7QisI1Iw20 12hZ14KKKX_9hFejX9PnygYOeY4KBs4Fb 1QEHtNeOq7LQBwaF9UFA7skE8X3bx7wDI 1Ln4vuWM9AMIOgtkDMM6MMNDp0NKXAg7c 1VmX36EMx2GSbGLO-zeu3_CY1-WA6fekU".split()

# ---------------- 建立文件 ----------------
doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(2.0); sec.right_margin = Cm(2.0)
style = doc.styles['Normal']
style.font.name = FONT
style.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)

# 標題
para(doc, "114 學年度「南科綠手指」教師專業學習社群成果報告",
     size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x1B, 0x5E, 0x20), space_after=4)
para(doc, "國立南科國際實驗高級中學（國小部）",
     size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

# 一、基本資料
add_heading_bar(doc, "壹、社群基本資料")
info_table(doc, [
    ("學校名稱", "國立南科國際實驗高級中學（國小部）", "社群名稱", "南科綠手指"),
    ("召集人", "陳賢宗", "Email", "prayer@ms.nnkieh.tn.edu.tw"),
    ("電話", "0988236863", "社群人數", "本校自然、藝文及跨領域教師"),
    ("社群目標",
     "一、精進教師在自然領域植物單元的教學知能。\n"
     "二、提升教師校園環境綠美化與香草、蔬果栽種的能力。\n"
     "三、豐富校園植物生物多樣性，建構食農與環境教育場域。\n"
     "四、提供教師與學生自然素材在藝術創作與生活的應用。\n"
     "五、推動環境教育，培訓環境教育種子教師。"),
    ("預期成效",
     "校園環境綠美化、自然素材跨領域應用、環境教育模組教學資源應用、"
     "辦理週三教師專業成長研習、辦理公開授課（含備觀議課）、"
     "自然與食農教學資源整合與教案開發。"),
    ("運作內容",
     "☑備觀議課（教師公開授課與專業回饋）　☑跨領域教學設計（協同教學）\n"
     "☑十二年國教議題—環境教育：香草、艾草、秋葵、橡實等植物之生活應用與栽種\n"
     "☑特定教育專業主題探索：校園綠美化、食農教育與數位互動教材應用"),
])

# 二、社群運作研習活動表
add_heading_bar(doc, "貳、社群運作活動一覽表")
t = doc.add_table(rows=1, cols=6)
t.style = 'Table Grid'
t.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = t.rows[0].cells
for i, h in enumerate(["場次", "日期", "活動名稱", "實施內容", "講師/主持人", "地點"]):
    cell_text(hdr[i], h, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill=HDR)
rows_data = [
    ("1", "114/09/24", "第一次社聚—香草植物的生活應用", "香草植物介紹與生活應用實作", "陳賢宗", "國小部自然教室一"),
    ("2", "114/11/19", "第二次社聚—香草植物的栽種", "香草扦插、育苗與栽種實作", "陳賢宗", "國小部自然教室一"),
    ("3", "114/12/24", "第三次社聚—認識可愛的橡實", "橡實種類認識與自然素材創作", "陳賢宗", "國小部自然教室一"),
    ("4", "115/02/25", "第四次社聚—艾草的生活應用與栽種", "艾草的生活應用與栽種實作", "陳賢宗", "國小部自然教室一"),
    ("5", "115/03", "第五次社聚—秋葵的生活應用與栽種", "秋葵育苗、栽種與飲食應用", "陳賢宗", "國小部自然教室一"),
    ("6", "線上", "第六次社聚—開心農場數位互動教材的應用", "線上會議：數位互動教材融入食農教學", "陳賢宗", "線上會議"),
    ("研習", "115/05/20", "昆蟲旅館研習", "昆蟲旅館設置與校園生態多樣性研習", "陳賢宗", "國小部自然教室一"),
    ("公開授課一", "114/11/07", "Rebecca 老師公開授課", "教師公開授課（含備、觀、議課）", "Rebecca 老師", "國小部教室"),
    ("公開授課二", "115/05/12", "公開授課", "教師公開授課（含備、觀、議課）", "本社群教師", "國小部教室"),
]
for rd in rows_data:
    cells = t.add_row().cells
    for i, v in enumerate(rd):
        cell_text(cells[i], v, size=11,
                  align=WD_ALIGN_PARAGRAPH.CENTER if i in (0,1,4,5) else WD_ALIGN_PARAGRAPH.LEFT)

# 三、社群運作紀錄與照片
add_heading_bar(doc, "參、社群運作紀錄與活動照片")

def activity(doc, title, meta, desc, ids, cols=3, width=5.0):
    para(doc, title, size=13, bold=True, color=RGBColor(0x1B,0x5E,0x20), space_after=2)
    if meta:
        para(doc, meta, size=11, space_after=2)
    if desc:
        para(doc, desc, size=11, space_after=4)
    photo_grid(doc, ids, cols=cols, width_cm=width)

activity(doc, "第一次社聚—香草植物的生活應用",
         "時間：114年9月24日　地點：國小部自然教室一",
         "介紹迷迭香、薄荷等常見香草植物的特性，並帶領教師體驗香草在生活與飲食中的應用，"
         "為後續校園栽種與跨領域教學奠定基礎。", A1)
activity(doc, "第二次社聚—香草植物的栽種",
         "時間：114年11月19日　地點：國小部自然教室一",
         "實作香草植物的扦插、育苗與栽種，教師親手操作並討論如何將栽種歷程融入自然與食農課程。", A2)
activity(doc, "第三次社聚—認識可愛的橡實",
         "時間：114年12月24日　地點：國小部自然教室一",
         "認識校園與周遭常見橡實（殼斗科）種類，並運用橡實等自然素材進行藝術創作，"
         "體會自然素材在生活與美感教育的價值。", A3)
activity(doc, "第四次社聚—艾草的生活應用與栽種",
         "時間：115年2月25日　地點：國小部自然教室一",
         "介紹艾草的植物特性與民俗、生活應用，並進行艾草栽種實作，"
         "連結節慶文化與環境教育。", A4)
activity(doc, "第五次社聚—秋葵的生活應用與栽種",
         "時間：115年3月　地點：國小部自然教室一",
         "以秋葵為主題，進行育苗與栽種實作，並探討秋葵在飲食與健康的應用，深化食農教育。", A5)
activity(doc, "第六次社聚—開心農場數位互動教材的應用（線上會議）",
         "形式：線上會議",
         "透過線上會議分享「開心農場」數位互動教材，討論如何將數位工具融入食農與自然教學，"
         "提升學生學習動機。", A6, cols=2, width=7.0)

# 四、研習
add_heading_bar(doc, "肆、自辦研習—昆蟲旅館研習")
activity(doc, "昆蟲旅館研習",
         "時間：115年5月20日　地點：國小部自然教室一",
         "帶領教師認識昆蟲旅館的功能與設置方式，透過實作提升校園生態多樣性，"
         "並培養教師將生態議題融入教學的能力。", R)

# 五、公開授課
add_heading_bar(doc, "伍、公開授課紀錄（備、觀、議課）")
activity(doc, "公開授課一—Rebecca 老師公開授課",
         "授課者：Rebecca 老師　時間：114年11月7日　地點：國小部教室",
         "本社群透過教師公開授課進行備課、觀課與議課，落實同儕專業對話，精進教學策略與課程設計。", P1)
activity(doc, "公開授課二",
         "時間：115年5月12日　地點：國小部教室",
         "延續社群專業成長，進行第二場公開授課與專業回饋，"
         "深化跨領域教學設計與教師學習共同體之建構。", P2)

# 結語
add_heading_bar(doc, "陸、結語")
para(doc,
     "本學年度「南科綠手指」教師專業學習社群以香草、橡實、艾草、秋葵等植物為主軸，"
     "結合六次社聚、昆蟲旅館研習及兩場公開授課，將自然、食農與環境教育融入教學現場。"
     "透過教師共同備課、觀課與議課，不僅提升教師專業知能與教學媒材研發能力，"
     "也豐富了校園植物與生態的多樣性，落實環境教育並回饋於學生學習。未來將持續深化社群運作，"
     "開發更多跨領域教學資源，並將成果與校內外教師分享。", size=12)

doc.save(OUT)
print("saved", OUT)
